#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 ysty2（英文文献）和 ysty3（知网文献）提取数据，生成 literature_full.json
"""
import json
import re
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_JSON = BASE_DIR / "data" / "literature_full.json"

YSTY2_DIR = BASE_DIR.parent / "ysty2" / "获得数据"
YSTY3_DIR = BASE_DIR.parent / "ysty3" / "获得数据"


def extract_herbs(text):
    """从涉及物质文本中提取物质名称列表"""
    if not text:
        return []
    text = str(text).strip()
    # 替换常见分隔符
    text = text.replace('；', ';').replace('，', ',').replace('、', ',')
    parts = [p.strip() for p in re.split(r'[;,]', text) if p.strip()]
    return parts


def truncate(text, max_len=200):
    if not text:
        return ""
    text = str(text).strip()
    if len(text) > max_len:
        return text[:max_len] + "…"
    return text


def clean_record(rec):
    """去除空值字段，减小体积"""
    return {k: v for k, v in rec.items() if v and v != "None" and v != "nan"}


def load_ysty2():
    """加载英文文献数据"""
    try:
        import openpyxl
    except ImportError:
        print("[ERROR] openpyxl not installed")
        return []

    files = [
        "食药物质英文文献数据库_2021-2026_第1部分.xlsx",
        "食药物质英文文献数据库_2021-2026_第2部分.xlsx",
        "食药物质英文文献数据库_2021-2026_第3部分.xlsx",
        "食药物质英文文献数据库_2021-2026_第4部分.xlsx",
    ]

    records = []
    seen_titles = set()

    for fname in files:
        fpath = YSTY2_DIR / fname
        if not fpath.exists():
            print(f"[WARN] 缺失文件: {fpath}")
            continue
        print(f"[LOAD] {fname}")
        wb = openpyxl.load_workbook(fpath, data_only=True)
        ws = wb["文献明细"]

        for row in range(2, ws.max_row + 1):
            title = ws.cell(row, 3).value
            if not title or title in seen_titles:
                continue
            seen_titles.add(title)

            herbs_text = ws.cell(row, 2).value
            herbs = extract_herbs(herbs_text)
            if not herbs:
                continue

            rec = {
                "id": len(records) + 1,
                "title": str(title).strip(),
                "authors": truncate(str(ws.cell(row, 4).value or ""), 100),
                "source": str(ws.cell(row, 6).value or "").strip(),
                "year": str(ws.cell(row, 7).value or "").strip(),
                "herbs": herbs,
                "type": "english",
                "category": "modern",
            }
            records.append(rec)

    print(f"[OK] ysty2 共加载 {len(records)} 条英文文献")
    return records


def load_ysty3():
    """加载知网文献数据"""
    try:
        import openpyxl
    except ImportError:
        print("[ERROR] openpyxl not installed")
        return []

    files = [
        "106种食药物质文献汇总-02-知网文献分册1（丁香至鸡内金）.xlsx",
        "106种食药物质文献汇总-03-知网文献分册2（麦芽至菊苣）.xlsx",
        "106种食药物质文献汇总-04-知网文献分册3（黄精至西洋参）.xlsx",
        "106种食药物质文献汇总-05-知网文献分册4（黄芪至化橘红）.xlsx",
    ]

    records = []
    seen_titles = set()

    for fname in files:
        fpath = YSTY3_DIR / fname
        if not fpath.exists():
            print(f"[WARN] 缺失文件: {fpath}")
            continue
        print(f"[LOAD] {fname}")
        wb = openpyxl.load_workbook(fpath, data_only=True)
        ws = wb["知网文献"]

        for row in range(2, ws.max_row + 1):
            title = ws.cell(row, 3).value
            if not title or title in seen_titles:
                continue
            seen_titles.add(title)

            herbs_text = ws.cell(row, 1).value
            herbs = extract_herbs(herbs_text)
            if not herbs:
                continue

            rec = {
                "id": len(records) + 1,
                "title": str(title).strip(),
                "authors": truncate(str(ws.cell(row, 4).value or ""), 100),
                "source": str(ws.cell(row, 2).value or "").strip(),
                "year": str(ws.cell(row, 10).value or "").strip(),
                "herbs": herbs,
                "type": "chinese",
                "category": "modern",
            }
            records.append(rec)

    print(f"[OK] ysty3 共加载 {len(records)} 条中文文献")
    return records


def main():
    english_records = load_ysty2()
    chinese_records = load_ysty3()

    all_records = english_records + chinese_records
    # 重新编号
    for i, rec in enumerate(all_records):
        rec["id"] = i + 1

    # 清理空字段
    all_records = [clean_record(r) for r in all_records]
    # 输出（紧凑格式，无缩进）
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(all_records, f, ensure_ascii=False, separators=(',', ':'))

    print(f"\n[OUTPUT] 共 {len(all_records)} 条文献，输出到 {OUTPUT_JSON}")
    print(f"  英文: {len(english_records)} 条")
    print(f"  中文: {len(chinese_records)} 条")

    # 文件大小估算
    import os
    size_mb = os.path.getsize(OUTPUT_JSON) / (1024 * 1024)
    print(f"  文件大小: {size_mb:.2f} MB")


if __name__ == "__main__":
    main()
