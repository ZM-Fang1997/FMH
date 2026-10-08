#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 website/data/substances/*.xlsx 批量提取药典字段，生成 substances_extracted.json
"""
import os
import re
import json
import openpyxl
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
XLSX_DIR = BASE_DIR / "data" / "substances"
OUTPUT_DIR = BASE_DIR / "data"
DATA_JS = BASE_DIR / "data.js"


def load_substances_meta():
    """从 data.js 读取 substances 数组"""
    with open(DATA_JS, "r", encoding="utf-8") as f:
        content = f.read()
    start = content.find("var substances = [")
    if start == -1:
        raise ValueError("Cannot find substances array in data.js")
    # 手动追踪括号深度找到匹配的 ]
    bracket_depth = 0
    in_string = False
    string_char = None
    i = start + len("var substances = [")
    while i < len(content):
        ch = content[i]
        if not in_string:
            if ch in ('"', "'"):
                in_string = True
                string_char = ch
            elif ch == '[':
                bracket_depth += 1
            elif ch == ']':
                if bracket_depth == 0:
                    end = i
                    break
                bracket_depth -= 1
        else:
            if ch == string_char and content[i-1] != '\\':
                in_string = False
        i += 1
    array_text = content[start + len("var substances = "):end+1]
    return json.loads(array_text)


def extract_latin_from_alias(text):
    """从别名字段提取拉丁名，如 'Syzygium aromaticum (L.) Merr. et Perry'"""
    if not text:
        return ""
    # 匹配属名 种加词 的拉丁学名模式
    # 通常拉丁名在中文别名之后，以属名大写开头
    m = re.search(r"([A-Z][a-z]+(?:\s+[a-z]+){1,2}(?:\s+\(?[A-Z][a-z\.]+\)?)?)", text)
    if m:
        candidate = m.group(1).strip()
        # 排除纯英文名（如 Clove, Star Anise）—— 英文名通常不含属+种结构
        if re.search(r"[a-z]+\s+[a-z]+", candidate):
            return candidate
    return ""


def extract_latin_from_origin(text):
    """从基原字段提取拉丁学名"""
    if not text:
        return ""
    m = re.search(r"([A-Z][a-z]+\s+[a-z]+(?:\s+[A-Z][a-z\.]+)?)", text)
    if m:
        return m.group(1).strip()
    return ""


def parse_nature_flavor_meridian(text):
    """从性味归经拆分四气、五味、归经"""
    if not text:
        return "", "", ""
    text = str(text)
    nature = ""
    flavor = ""
    meridians = ""

    # 四气通常在句尾或逗号后
    nature_match = re.search(r"[，。,\.\s]([寒热温凉平]+)[。\.\s]", text)
    if nature_match:
        nature = nature_match.group(1)
    else:
        nature_match = re.search(r"([寒热温凉平]+)[。\.\s]", text)
        if nature_match:
            nature = nature_match.group(1)

    # 五味：在"性味"之前，通常是"辛，温"中的"辛"
    # 先尝试匹配 "X，Y" 模式
    nf_match = re.match(r"([辛甘酸苦咸淡涩、]+)[，,]\s*([寒热温凉平]+)", text)
    if nf_match:
        flavor = nf_match.group(1)
        if not nature:
            nature = nf_match.group(2)
    else:
        #  fallback：查找所有五味字符
        flavor_chars = re.findall(r"[辛甘酸苦咸淡涩]", text)
        if flavor_chars:
            flavor = "、".join(flavor_chars[:4])

    # 归经
    meridian_match = re.search(r"归([一-龥、]+?)(?:经|脉)", text)
    if meridian_match:
        meridians = meridian_match.group(1) + "经"

    return nature, flavor, meridians


def extract_dosage_range(text):
    """从用法用量提取剂量范围"""
    if not text:
        return None
    text = str(text)
    # 匹配 "3～6g" "3~6g" "3-6g" "3—6g"
    m = re.search(r"(\d+(?:\.\d+)?)\s*[～~\-—]\s*(\d+(?:\.\d+)?)\s*(g|克|mg|枚|个|片|丸|ml|毫升|条|只)", text)
    if m:
        return {
            "min": float(m.group(1)),
            "max": float(m.group(2)),
            "unit": m.group(3),
            "raw": m.group(0)
        }
    # 匹配单个数字+单位
    m2 = re.search(r"(\d+(?:\.\d+)?)\s*(g|克|mg|枚|个|片|丸|ml|毫升|条|只)", text)
    if m2:
        v = float(m2.group(1))
        return {
            "min": v,
            "max": v,
            "unit": m2.group(2),
            "raw": m2.group(0)
        }
    return None


def extract_fields_from_xlsx(filepath):
    """读取单个 xlsx，提取关键字段"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    fields = {}
    # 遍历前 50 行提取标签-值对
    for row in range(1, min(51, ws.max_row + 1)):
        label = ws.cell(row, 1).value
        value = ws.cell(row, 2).value
        if label and isinstance(label, str):
            label = label.strip()
            if value and isinstance(value, str):
                value = value.strip()
            fields[label] = value
    return fields


def build_substance_record(meta, fields):
    """整合元数据和提取字段"""
    name = meta["name"]

    # 拉丁名：优先从别名（英文名、拉丁名）提取，其次基原，最后 fallback 到 en_name
    alias_text = fields.get("别名（英文名、拉丁名）", "")
    latin = extract_latin_from_alias(alias_text)
    if not latin:
        latin = extract_latin_from_origin(fields.get("基原", ""))
    if not latin:
        en = meta.get("en_name", "")
        # 如果 en_name 本身像拉丁名（两个单词，首字母大写+小写）
        if re.match(r"^[A-Z][a-z]+\s+[a-z]+", en):
            latin = en
        else:
            latin = ""  # 留空，后续统一补充

    nature, flavor, meridians = parse_nature_flavor_meridian(fields.get("性味归经", ""))
    dosage_range = extract_dosage_range(fields.get("用法用量", ""))

    record = {
        "id": meta.get("id"),
        "name": name,
        "latin_name": latin,
        "en_name": meta.get("en_name", ""),
        "image": meta.get("image", ""),
        "source_status": meta.get("source_status", ""),
        "origin": fields.get("基原", ""),
        "description": fields.get("性状", ""),
        "identification": fields.get("鉴别", ""),
        "inspection": fields.get("检查", ""),
        "assay": fields.get("含量测定", ""),
        "processing": fields.get("炮制", ""),
        "nature": nature,
        "flavor": flavor,
        "meridians": meridians,
        "functions": fields.get("功能主治", ""),
        "usage_dosage": fields.get("用法用量", ""),
        "dosage_range": dosage_range,
        "contraindications": fields.get("禁忌症", "") or fields.get("注意事项", ""),
        "toxicity": fields.get("毒性", "") or fields.get("副作用", ""),
        "storage": fields.get("贮藏", ""),
        "major_components": fields.get("主要成分", ""),
        "pharmacology": fields.get("药理作用", ""),
        "suitable_population": fields.get("宜忌人群", "") or fields.get("适宜人群", ""),
        "dietary_taboo": fields.get("食养禁忌", ""),
        "dietary_efficacy": fields.get("食疗功效", "") or fields.get("食养功效", ""),
        "english_name": fields.get("英文名称", "") or fields.get("英文名", ""),
        "alias": alias_text,
        "producing_area": fields.get("主产地及科属划分", "") or fields.get("产地", ""),
        "key_compounds": fields.get("关键化合物", ""),
        "part": fields.get("部位", ""),
        "source": fields.get("来源", ""),
        "medicinal_property": fields.get("药性", ""),
        "usage_advice": fields.get("食用建议", ""),
    }
    return record


def main():
    substances_meta = load_substances_meta()
    results = []

    for meta in substances_meta:
        name = meta["name"]
        xlsx_path = XLSX_DIR / f"{name}.xlsx"
        if not xlsx_path.exists():
            print(f"[WARN] 缺失文件: {xlsx_path}")
            record = {
                "id": meta.get("id"),
                "name": name,
                "latin_name": meta.get("en_name", ""),
                "en_name": meta.get("en_name", ""),
                "image": meta.get("image", ""),
                "source_status": meta.get("source_status", ""),
            }
            results.append(record)
            continue

        fields = extract_fields_from_xlsx(xlsx_path)
        record = build_substance_record(meta, fields)
        results.append(record)
        dr = record['dosage_range']
        dr_str = f"{dr['min']}~{dr['max']}{dr['unit']}" if dr else "无"
        print(f"[OK] id={record['id']:3} {name}: 性味={record['nature']}{record['flavor']}, 归经={record['meridians']}, 剂量={dr_str}, 拉丁={record['latin_name'][:30] if record['latin_name'] else '空'}")

    output_path = OUTPUT_DIR / "substances_extracted.json"
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\n提取完成，共 {len(results)} 条，输出到 {output_path}")


if __name__ == "__main__":
    main()
